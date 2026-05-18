import argparse
import os
import sys
import time
from contextlib import contextmanager, nullcontext
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import torch
import torch.nn.functional as F
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font
from tqdm import tqdm


SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

WTConvLNAMDStatic = None
MSM = None
AnomalyMapOptimizer = None
RsCIN = None
open_clip = None
_backbones = None
mvtec = None
visa = None
btad = None
miniled = None
microled = None
_CLASSNAMES_mvtec_ad = None
_CLASSNAMES_visa = None
_CLASSNAMES_btad = None
_CLASSNAMES_miniled = None
_CLASSNAMES_microled = None


def load_project_modules():
    global WTConvLNAMDStatic
    global MSM
    global AnomalyMapOptimizer
    global RsCIN
    global open_clip
    global _backbones
    global mvtec
    global visa
    global btad
    global miniled
    global microled
    global _CLASSNAMES_mvtec_ad
    global _CLASSNAMES_visa
    global _CLASSNAMES_btad
    global _CLASSNAMES_miniled
    global _CLASSNAMES_microled

    import datasets.mvtec as _mvtec
    import datasets.visa as _visa
    import datasets.btad as _btad
    import datasets.miniled as _miniled
    import datasets.microled as _microled
    import models.backbone._backbones as __backbones
    import models.backbone.open_clip as _open_clip
    from datasets.btad import _CLASSNAMES as __CLASSNAMES_btad
    from datasets.microled import _CLASSNAMES as __CLASSNAMES_microled
    from datasets.miniled import _CLASSNAMES as __CLASSNAMES_miniled
    from datasets.mvtec import _CLASSNAMES as __CLASSNAMES_mvtec_ad
    from datasets.visa import _CLASSNAMES as __CLASSNAMES_visa
    from models.modules.WTConvStatic import WTConvLNAMDStatic as _WTConvLNAMDStatic
    from models.modules._MSM import MSM as _MSM
    from models.modules._Optimization import AnomalyMapOptimizer as _AnomalyMapOptimizer
    from models.modules._RsCIN import RsCIN as _RsCIN

    WTConvLNAMDStatic = _WTConvLNAMDStatic
    MSM = _MSM
    AnomalyMapOptimizer = _AnomalyMapOptimizer
    RsCIN = _RsCIN
    open_clip = _open_clip
    _backbones = __backbones
    mvtec = _mvtec
    visa = _visa
    btad = _btad
    miniled = _miniled
    microled = _microled
    _CLASSNAMES_mvtec_ad = __CLASSNAMES_mvtec_ad
    _CLASSNAMES_visa = __CLASSNAMES_visa
    _CLASSNAMES_btad = __CLASSNAMES_btad
    _CLASSNAMES_miniled = __CLASSNAMES_miniled
    _CLASSNAMES_microled = __CLASSNAMES_microled


@dataclass
class ProfileRecord:
    category: str
    divide_iter: str
    module: str
    images: int
    total_ms: float
    ms_per_image: float
    start_allocated_mb: float
    end_allocated_mb: float
    peak_allocated_mb: float
    peak_extra_allocated_mb: float
    start_reserved_mb: float
    end_reserved_mb: float
    peak_reserved_mb: float
    peak_extra_reserved_mb: float


def resolve_existing_path(path_value):
    path = Path(path_value)
    candidates = [
        path if path.is_absolute() else Path.cwd() / path,
        path if path.is_absolute() else REPO_ROOT / path,
        path if path.is_absolute() else SCRIPT_DIR / path,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()
    return candidates[0].resolve()


def resolve_output_path(path_value):
    path = Path(path_value)
    if path.is_absolute():
        return path
    return (REPO_ROOT / path).resolve()


def parse_bool(value):
    if value is None:
        return None
    return str(value).lower() in ("1", "true", "yes", "y")


def load_yaml_config(config_path):
    config_path = resolve_existing_path(config_path)
    with open(config_path, "r", encoding="utf-8") as f:
        return yaml.load(f, Loader=yaml.FullLoader), config_path


def get_args():
    parser = argparse.ArgumentParser(description="Profile MuSc module time and CUDA memory.")
    parser.add_argument("--config", type=str, default="configs/musc.yaml", help="config file path")
    parser.add_argument("--data_path", type=str, default=None, help="dataset path")
    parser.add_argument("--dataset_name", type=str, default=None, help="dataset name")
    parser.add_argument("--class_name", type=str, default=None, help="category")
    parser.add_argument("--device", type=int, default=None, help="gpu id")
    parser.add_argument("--output_dir", type=str, default=None, help="save results path")
    parser.add_argument("--output_excel", type=str, default=None, help="profile xlsx output path")
    parser.add_argument("--output_txt", type=str, default=None, help="legacy profile output path; saved as xlsx")
    parser.add_argument("--vis", type=str, default=None, help="visualization")
    parser.add_argument("--vis_type", type=str, default=None, help="normalization type in visualization")
    parser.add_argument("--save_excel", type=str, default=None, help="save excel")
    parser.add_argument("--r_list", type=int, nargs="+", default=None, help="aggregation degrees of LNAMD")
    parser.add_argument("--feature_layers", type=int, nargs="+", default=None, help="feature layers")
    parser.add_argument("--backbone_name", type=str, default=None, help="backbone")
    parser.add_argument("--pretrained", type=str, default=None, help="pretrained datasets")
    parser.add_argument("--img_resize", type=int, default=None, help="image size")
    parser.add_argument("--batch_size", type=int, default=None, help="batch size")
    parser.add_argument("--divide_num", type=int, default=None, help="the number of divided subsets")
    parser.add_argument(
        "--max_images",
        type=int,
        default=None,
        help="optional image limit per category for a quick profile; MSM needs at least 5 images",
    )
    return parser.parse_args()


def apply_args(cfg, args):
    if args.data_path is not None:
        cfg["datasets"]["data_path"] = args.data_path
    data_path = Path(cfg["datasets"]["data_path"])
    if not data_path.exists():
        raise FileNotFoundError(f"The dataset path {data_path} does not exist.")

    if args.dataset_name is not None:
        cfg["datasets"]["dataset_name"] = args.dataset_name
    if args.class_name is not None:
        cfg["datasets"]["class_name"] = args.class_name
    if args.device is not None:
        cfg["device"] = args.device
    if isinstance(cfg["device"], int):
        cfg["device"] = str(cfg["device"])
    if args.output_dir is not None:
        cfg["testing"]["output_dir"] = args.output_dir
    os.makedirs(cfg["testing"]["output_dir"], exist_ok=True)

    vis = parse_bool(args.vis)
    if vis is not None:
        cfg["testing"]["vis"] = vis
    if args.vis_type is not None:
        cfg["testing"]["vis_type"] = args.vis_type
    save_excel = parse_bool(args.save_excel)
    if save_excel is not None:
        cfg["testing"]["save_excel"] = save_excel
    if args.r_list is not None:
        cfg["models"]["r_list"] = args.r_list
    if isinstance(cfg["models"]["r_list"], int):
        cfg["models"]["r_list"] = [cfg["models"]["r_list"]]
    if args.feature_layers is not None:
        cfg["models"]["feature_layers"] = args.feature_layers
    if isinstance(cfg["models"]["feature_layers"], int):
        cfg["models"]["feature_layers"] = [cfg["models"]["feature_layers"]]
    if args.backbone_name is not None:
        cfg["models"]["backbone_name"] = args.backbone_name
    if args.pretrained is not None:
        cfg["models"]["pretrained"] = args.pretrained
    if args.img_resize is not None:
        cfg["datasets"]["img_resize"] = args.img_resize
    if args.batch_size is not None:
        cfg["models"]["batch_size"] = args.batch_size
    if args.divide_num is not None:
        cfg["datasets"]["divide_num"] = args.divide_num
    return cfg


class MuScProfileRunner:
    def __init__(self, cfg, seed=0):
        self.cfg = cfg
        self.seed = seed
        self.device = torch.device(
            f"cuda:{cfg['device']}" if torch.cuda.is_available() else "cpu"
        )
        print(f"Active device: {self.device}")

        self.path = cfg["datasets"]["data_path"]
        self.dataset = cfg["datasets"]["dataset_name"]
        self.categories = cfg["datasets"]["class_name"]
        if isinstance(self.categories, str):
            if self.categories.lower() == "all":
                if self.dataset == "visa":
                    self.categories = _CLASSNAMES_visa
                elif self.dataset == "mvtec_ad":
                    self.categories = _CLASSNAMES_mvtec_ad
                elif self.dataset == "btad":
                    self.categories = _CLASSNAMES_btad
                elif self.dataset == "miniled_ad":
                    self.categories = _CLASSNAMES_miniled
                elif self.dataset == "microled_ad":
                    self.categories = _CLASSNAMES_microled
            else:
                self.categories = [self.categories]

        self.model_name = cfg["models"]["backbone_name"]
        self.image_size = cfg["datasets"]["img_resize"]
        self.batch_size = cfg["models"]["batch_size"]
        self.pretrained = cfg["models"]["pretrained"]
        self.features_list = [l + 1 for l in cfg["models"]["feature_layers"]]
        self.divide_num = cfg["datasets"]["divide_num"]
        self.r_list = cfg["models"]["r_list"]
        self.output_dir = os.path.join(
            cfg["testing"]["output_dir"],
            self.dataset,
            self.model_name,
            f"imagesize{self.image_size}",
        )
        os.makedirs(self.output_dir, exist_ok=True)
        self.load_backbone()

    def load_backbone(self):
        if "dino" in self.model_name:
            self.dino_model = _backbones.load(self.model_name)
            self.dino_model.to(self.device)
            self.dino_model.eval()
            self.preprocess = None
        else:
            self.clip_model, _, self.preprocess = open_clip.create_model_and_transforms(
                self.model_name,
                self.image_size,
                pretrained=self.pretrained,
            )
            self.clip_model.to(self.device)
            self.clip_model.eval()

    def load_datasets(self, category, divide_num=1, divide_iter=0):
        common_args = {
            "source": self.path,
            "classname": category,
            "resize": self.image_size,
            "imagesize": self.image_size,
            "clip_transformer": self.preprocess,
            "divide_num": divide_num,
            "divide_iter": divide_iter,
            "random_seed": self.seed,
        }
        if self.dataset == "visa":
            return visa.VisaDataset(split=visa.DatasetSplit.TEST, **common_args)
        if self.dataset == "mvtec_ad":
            return mvtec.MVTecDataset(split=mvtec.DatasetSplit.TEST, **common_args)
        if self.dataset == "btad":
            return btad.BTADDataset(split=btad.DatasetSplit.TEST, **common_args)
        if self.dataset == "miniled_ad":
            return miniled.MiniledDataset(split=miniled.DatasetSplit.TEST, **common_args)
        if self.dataset == "microled_ad":
            return microled.MicroledDataset(split=microled.DatasetSplit.TEST, **common_args)
        raise ValueError(f"Unsupported dataset_name: {self.dataset}")


class ModuleProfiler:
    def __init__(self, device):
        self.device = device
        self.use_cuda = device.type == "cuda" and torch.cuda.is_available()
        self.records = []

    def _sync(self):
        if self.use_cuda:
            torch.cuda.synchronize(self.device)

    def _memory_mb(self):
        if not self.use_cuda:
            return 0.0, 0.0
        allocated = torch.cuda.memory_allocated(self.device) / 1024 / 1024
        reserved = torch.cuda.memory_reserved(self.device) / 1024 / 1024
        return allocated, reserved

    def _peak_memory_mb(self):
        if not self.use_cuda:
            return 0.0, 0.0
        allocated = torch.cuda.max_memory_allocated(self.device) / 1024 / 1024
        reserved = torch.cuda.max_memory_reserved(self.device) / 1024 / 1024
        return allocated, reserved

    @contextmanager
    def measure(self, category, divide_iter, module, images):
        images = max(int(images), 1)
        self._sync()
        start_allocated, start_reserved = self._memory_mb()
        if self.use_cuda:
            torch.cuda.reset_peak_memory_stats(self.device)
        start = time.perf_counter()
        yield
        self._sync()
        total_ms = (time.perf_counter() - start) * 1000
        end_allocated, end_reserved = self._memory_mb()
        peak_allocated, peak_reserved = self._peak_memory_mb()
        self.records.append(
            ProfileRecord(
                category=str(category),
                divide_iter=str(divide_iter),
                module=str(module),
                images=images,
                total_ms=total_ms,
                ms_per_image=total_ms / images,
                start_allocated_mb=start_allocated,
                end_allocated_mb=end_allocated,
                peak_allocated_mb=peak_allocated,
                peak_extra_allocated_mb=max(0.0, peak_allocated - start_allocated),
                start_reserved_mb=start_reserved,
                end_reserved_mb=end_reserved,
                peak_reserved_mb=peak_reserved,
                peak_extra_reserved_mb=max(0.0, peak_reserved - start_reserved),
            )
        )


def autocast_context(device):
    if device.type == "cuda" and torch.cuda.is_available():
        return torch.cuda.amp.autocast()
    return nullcontext()


def empty_cuda_cache(device):
    if device.type == "cuda" and torch.cuda.is_available():
        torch.cuda.empty_cache()


def limit_dataset(test_dataset, max_images):
    if max_images is None:
        return
    if max_images < 5:
        raise ValueError("MSM uses mutual scoring and needs at least 5 images when --max_images is set.")
    if hasattr(test_dataset, "data_to_iterate"):
        test_dataset.data_to_iterate = test_dataset.data_to_iterate[:max_images]
    else:
        raise AttributeError("This dataset does not expose data_to_iterate, so --max_images cannot be applied.")


def ablation_config():
    return {
        "wt_type": "db1",
        "padding": "reflect",
        "level0": False,
        "use_details": True,
        "detail_start": 1,
        "keep_ll": True,
        "gamma": 2.0,
        "use_spot_weight": True,
        "use_morphology": True,
        "morph_open_k": 1,
        "morph_close_k": 3,
        "morph_smooth_k": 3,
        "morph_sigma": 0.5,
    }


def extract_features(model, test_dataloader, image_count, category, divide_iter, profiler):
    patch_tokens_list = []
    class_tokens = []

    module_name = "BackboneFeatureExtraction"
    with profiler.measure(category, divide_iter, module_name, image_count):
        for image_info in tqdm(test_dataloader, desc=f"{category} feature extraction"):
            image = image_info["image"]

            with torch.no_grad(), autocast_context(model.device):
                input_image = image.to(torch.float).to(model.device)
                if "dinov2" in model.model_name or "dinov3" in model.model_name:
                    patch_tokens = model.dino_model.get_intermediate_layers(
                        x=input_image,
                        n=[l - 1 for l in model.features_list],
                        return_class_token=False,
                    )
                    image_features = model.dino_model(input_image)
                    patch_tokens = [patch_tokens[l].cpu() for l in range(len(model.features_list))]
                    fake_cls = [torch.zeros_like(p)[:, 0:1, :] for p in patch_tokens]
                    patch_tokens = [
                        torch.cat([fake_cls[i], patch_tokens[i]], dim=1)
                        for i in range(len(patch_tokens))
                    ]
                elif "dino" in model.model_name:
                    patch_tokens_all = model.dino_model.get_intermediate_layers(
                        x=input_image,
                        n=max(model.features_list),
                    )
                    image_features = model.dino_model(input_image)
                    patch_tokens = [patch_tokens_all[l - 1].cpu() for l in model.features_list]
                else:
                    image_features, patch_tokens = model.clip_model.encode_image(
                        input_image,
                        model.features_list,
                    )
                    image_features /= image_features.norm(dim=-1, keepdim=True)
                    patch_tokens = [patch_tokens[l].cpu() for l in range(len(model.features_list))]

            image_features = [
                image_features[bi].squeeze().cpu().numpy()
                for bi in range(image_features.shape[0])
            ]
            class_tokens.extend(image_features)
            patch_tokens_list.append(patch_tokens)

    return patch_tokens_list, class_tokens


def profile_category(model, category, profiler, max_images=None):
    cfg = ablation_config()
    anomaly_maps = torch.tensor([]).double()
    class_tokens_all = []
    total_images = 0

    for divide_iter in range(model.divide_num):
        test_dataset = model.load_datasets(
            category,
            divide_num=model.divide_num,
            divide_iter=divide_iter,
        )
        limit_dataset(test_dataset, max_images)
        subset_num = len(test_dataset)
        if subset_num < 5:
            raise ValueError(
                f"Category {category}, divide {divide_iter} has {subset_num} images. "
                "MSM needs at least 5 images."
            )
        total_images += subset_num

        test_dataloader = torch.utils.data.DataLoader(
            test_dataset,
            batch_size=model.batch_size,
            shuffle=False,
            num_workers=0,
            pin_memory=True,
        )

        patch_tokens_list, class_tokens = extract_features(
            model,
            test_dataloader,
            subset_num,
            category,
            divide_iter,
            profiler,
        )

        class_tokens_all.extend(class_tokens)

        feature_dim = patch_tokens_list[0][0].shape[-1]
        anomaly_maps_r = torch.tensor([]).double()

        for r in model.r_list:
            z_layers = {}
            with profiler.measure(category, divide_iter, f"WTConvLNAMDStatic-r{r}", subset_num):
                lnamd_r = WTConvLNAMDStatic(
                    device=model.device,
                    feature_dim=feature_dim,
                    feature_layer=model.features_list,
                    r=r,
                    wt_type=cfg["wt_type"],
                    padding_mode=cfg["padding"],
                    include_level0=cfg["level0"],
                    use_details=cfg["use_details"],
                    detail_start_level=cfg["detail_start"],
                    keep_ll=cfg["keep_ll"],
                )
                for patch_tokens_cpu in patch_tokens_list:
                    patch_tokens = [p.to(model.device) for p in patch_tokens_cpu]
                    with torch.no_grad(), autocast_context(model.device):
                        features = lnamd_r._embed(patch_tokens)
                        features /= features.norm(dim=-1, keepdim=True)
                        for layer_idx in range(len(model.features_list)):
                            key = str(layer_idx)
                            if key not in z_layers:
                                z_layers[key] = []
                            z_layers[key].append(features[:, :, layer_idx, :])
                    del patch_tokens

            anomaly_maps_l = torch.tensor([]).double()
            with profiler.measure(category, divide_iter, f"MSM-r{r}", subset_num):
                for layer_key in z_layers.keys():
                    z = torch.cat(z_layers[layer_key], dim=0).to(model.device)
                    current_use_spot_weight = cfg["use_spot_weight"]
                    if current_use_spot_weight:
                        is_target_category = (
                            model.dataset == "mvtec_ad"
                            and category in ["screw", "toothbrush", "zipper"]
                        )
                        if not is_target_category:
                            current_use_spot_weight = False

                    anomaly_maps_msm = MSM(
                        Z=z,
                        device=model.device,
                        topmin_min=0,
                        topmin_max=0.3,
                        gamma=cfg["gamma"],
                        use_spot_weight=current_use_spot_weight,
                    )
                    anomaly_maps_l = torch.cat(
                        (anomaly_maps_l, anomaly_maps_msm.unsqueeze(0).cpu()),
                        dim=0,
                    )
                    del z, anomaly_maps_msm
                    empty_cuda_cache(model.device)

                anomaly_maps_l = torch.mean(anomaly_maps_l, 0)
                anomaly_maps_r = torch.cat((anomaly_maps_r, anomaly_maps_l.unsqueeze(0)), dim=0)

            del z_layers
            empty_cuda_cache(model.device)

        with profiler.measure(category, divide_iter, "AverageRAndReshape", subset_num):
            anomaly_maps_iter = torch.mean(anomaly_maps_r, 0).to(model.device)
            batch_count, patch_count = anomaly_maps_iter.shape
            height = int(np.sqrt(patch_count))
            width = height
            anomaly_maps_iter_spatial = anomaly_maps_iter.view(batch_count, 1, height, width)

        if cfg["use_morphology"]:
            with profiler.measure(category, divide_iter, "AnomalyMapOptimizer", subset_num):
                optimizer = AnomalyMapOptimizer(
                    kernel_size_open=cfg["morph_open_k"],
                    kernel_size_close=cfg["morph_close_k"],
                    smooth_kernel_size=cfg["morph_smooth_k"],
                    sigma=cfg["morph_sigma"],
                ).to(model.device)
                anomaly_maps_iter_spatial = optimizer(anomaly_maps_iter_spatial)

        with profiler.measure(category, divide_iter, "InterpolateAndCollect", subset_num):
            anomaly_maps_iter = F.interpolate(
                anomaly_maps_iter_spatial,
                size=model.image_size,
                mode="bilinear",
                align_corners=True,
            )
            anomaly_maps = torch.cat((anomaly_maps, anomaly_maps_iter.squeeze(1).cpu()), dim=0)

        del anomaly_maps_r, anomaly_maps_iter, anomaly_maps_iter_spatial
        empty_cuda_cache(model.device)

    with profiler.measure(category, "all", "PrepareScores", total_images):
        anomaly_maps_np = anomaly_maps.cpu().numpy()
        image_count = anomaly_maps_np.shape[0]
        ac_score = np.array(anomaly_maps_np).reshape(image_count, -1).max(-1)

    with profiler.measure(category, "all", "RsCIN", total_images):
        if model.dataset == "visa":
            k_score = [1, 8, 9]
        elif model.dataset == "mvtec_ad":
            k_score = [1, 2, 3]
        else:
            k_score = [1, 2, 3]
        scores_cls = RsCIN(ac_score, class_tokens_all, k_list=k_score)

    return {
        "category": category,
        "images": total_images,
    }


def format_record(record):
    fields = [
        record.category,
        record.divide_iter,
        record.module,
        str(record.images),
        f"{record.total_ms:.3f}",
        f"{record.ms_per_image:.3f}",
        f"{record.start_allocated_mb:.2f}",
        f"{record.end_allocated_mb:.2f}",
        f"{record.peak_allocated_mb:.2f}",
        f"{record.peak_extra_allocated_mb:.2f}",
        f"{record.start_reserved_mb:.2f}",
        f"{record.end_reserved_mb:.2f}",
        f"{record.peak_reserved_mb:.2f}",
        f"{record.peak_extra_reserved_mb:.2f}",
    ]
    return "\t".join(fields)


def summarize_records(records):
    summary = {}
    for record in records:
        key = record.module
        if key not in summary:
            summary[key] = {
                "total_ms": 0.0,
                "images": 0,
                "max_peak_allocated": 0.0,
                "max_peak_reserved": 0.0,
                "max_peak_extra_allocated": 0.0,
                "max_peak_extra_reserved": 0.0,
            }
        summary[key]["total_ms"] += record.total_ms
        summary[key]["images"] += record.images
        summary[key]["max_peak_allocated"] = max(
            summary[key]["max_peak_allocated"],
            record.peak_allocated_mb,
        )
        summary[key]["max_peak_reserved"] = max(
            summary[key]["max_peak_reserved"],
            record.peak_reserved_mb,
        )
        summary[key]["max_peak_extra_allocated"] = max(
            summary[key]["max_peak_extra_allocated"],
            record.peak_extra_allocated_mb,
        )
        summary[key]["max_peak_extra_reserved"] = max(
            summary[key]["max_peak_extra_reserved"],
            record.peak_extra_reserved_mb,
        )
    return summary


def summarize_overall_profile(records, category_results, total_ms):
    total_images = sum(result["images"] for result in category_results)
    return {
        "images": total_images,
        "total_ms": total_ms,
        "weighted_ms_per_image": total_ms / max(total_images, 1),
        "max_peak_allocated": max(
            (record.peak_allocated_mb for record in records),
            default=0.0,
        ),
        "max_peak_extra_allocated": max(
            (record.peak_extra_allocated_mb for record in records),
            default=0.0,
        ),
        "max_peak_reserved": max(
            (record.peak_reserved_mb for record in records),
            default=0.0,
        ),
        "max_peak_extra_reserved": max(
            (record.peak_extra_reserved_mb for record in records),
            default=0.0,
        ),
    }


def summarize_combined_modules(summary, module_names, category_results):
    values = [summary[name] for name in module_names if name in summary]
    if not values:
        return None

    total_images = sum(result["images"] for result in category_results)
    total_ms = sum(value["total_ms"] for value in values)
    return {
        "images": total_images,
        "total_ms": total_ms,
        "weighted_ms_per_image": total_ms / max(total_images, 1),
        "max_peak_allocated": max(value["max_peak_allocated"] for value in values),
        "max_peak_extra_allocated": max(value["max_peak_extra_allocated"] for value in values),
        "max_peak_reserved": max(value["max_peak_reserved"] for value in values),
        "max_peak_extra_reserved": max(value["max_peak_extra_reserved"] for value in values),
    }


def profile_record_header():
    return [
        "category",
        "divide_iter",
        "module",
        "images",
        "total_ms",
        "ms_per_image",
        "start_allocated_mb",
        "end_allocated_mb",
        "peak_allocated_mb",
        "peak_extra_allocated_mb",
        "start_reserved_mb",
        "end_reserved_mb",
        "peak_reserved_mb",
        "peak_extra_reserved_mb",
    ]


def profile_record_row(record):
    return [
        record.category,
        record.divide_iter,
        record.module,
        record.images,
        record.total_ms,
        record.ms_per_image,
        record.start_allocated_mb,
        record.end_allocated_mb,
        record.peak_allocated_mb,
        record.peak_extra_allocated_mb,
        record.start_reserved_mb,
        record.end_reserved_mb,
        record.peak_reserved_mb,
        record.peak_extra_reserved_mb,
    ]


def normalize_excel_path(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    return path


def append_header(sheet, values):
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def autosize_columns(sheet):
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[column_letter].width = min(max_length + 2, 60)


def write_excel(output_excel, cfg, config_path, profiler, category_results, total_ms):
    output_excel = normalize_excel_path(output_excel)
    output_excel.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()

    config_sheet = workbook.active
    config_sheet.title = "RunConfig"
    append_header(config_sheet, ["item", "value"])
    config_rows = [
        ("config", str(config_path)),
        ("dataset", cfg["datasets"]["dataset_name"]),
        ("data_path", cfg["datasets"]["data_path"]),
        ("backbone", cfg["models"]["backbone_name"]),
        ("image_size", cfg["datasets"]["img_resize"]),
        ("batch_size", cfg["models"]["batch_size"]),
        ("r_list", str(cfg["models"]["r_list"])),
        ("feature_layers", str(cfg["models"]["feature_layers"])),
        ("total_wall_ms", total_ms),
    ]
    for row in config_rows:
        config_sheet.append(row)
    autosize_columns(config_sheet)

    category_sheet = workbook.create_sheet("CategoryProfile")
    append_header(category_sheet, ["category", "images"])
    for result in category_results:
        category_sheet.append([result["category"], result["images"]])
    autosize_columns(category_sheet)

    records_sheet = workbook.create_sheet("ModuleRecords")
    append_header(records_sheet, profile_record_header())
    for record in profiler.records:
        records_sheet.append(profile_record_row(record))
    autosize_columns(records_sheet)

    summary_sheet = workbook.create_sheet("SummaryByModule")
    append_header(
        summary_sheet,
        [
            "module",
            "images",
            "total_ms",
            "weighted_ms_per_image",
            "max_peak_allocated_mb",
            "max_peak_extra_allocated_mb",
            "max_peak_reserved_mb",
            "max_peak_extra_reserved_mb",
        ],
    )
    summary = summarize_records(profiler.records)
    for module, values in summary.items():
        images = max(values["images"], 1)
        summary_sheet.append(
            [
                module,
                values["images"],
                values["total_ms"],
                values["total_ms"] / images,
                values["max_peak_allocated"],
                values["max_peak_extra_allocated"],
                values["max_peak_reserved"],
                values["max_peak_extra_reserved"],
            ]
        )
    combined_wtconv = summarize_combined_modules(
        summary,
        ["WTConvLNAMDStatic-r1", "WTConvLNAMDStatic-r3", "WTConvLNAMDStatic-r5"],
        category_results,
    )
    if combined_wtconv is not None:
        summary_sheet.append(
            [
                "WTConvLNAMDStatic-r1+r3+r5",
                combined_wtconv["images"],
                combined_wtconv["total_ms"],
                combined_wtconv["weighted_ms_per_image"],
                combined_wtconv["max_peak_allocated"],
                combined_wtconv["max_peak_extra_allocated"],
                combined_wtconv["max_peak_reserved"],
                combined_wtconv["max_peak_extra_reserved"],
            ]
        )
    combined_msm = summarize_combined_modules(
        summary,
        ["MSM-r1", "MSM-r3", "MSM-r5"],
        category_results,
    )
    if combined_msm is not None:
        summary_sheet.append(
            [
                "MSM-r1+r3+r5",
                combined_msm["images"],
                combined_msm["total_ms"],
                combined_msm["weighted_ms_per_image"],
                combined_msm["max_peak_allocated"],
                combined_msm["max_peak_extra_allocated"],
                combined_msm["max_peak_reserved"],
                combined_msm["max_peak_extra_reserved"],
            ]
        )
    overall = summarize_overall_profile(profiler.records, category_results, total_ms)
    summary_sheet.append(
        [
            "MuScOverall",
            overall["images"],
            overall["total_ms"],
            overall["weighted_ms_per_image"],
            overall["max_peak_allocated"],
            overall["max_peak_extra_allocated"],
            overall["max_peak_reserved"],
            overall["max_peak_extra_reserved"],
        ]
    )
    autosize_columns(summary_sheet)

    workbook.save(output_excel)
    return output_excel


def main():
    args = get_args()
    load_project_modules()
    cfg, config_path = load_yaml_config(args.config)
    cfg = apply_args(cfg, args)

    output_path_arg = args.output_excel if args.output_excel is not None else args.output_txt
    if output_path_arg is None:
        output_excel = resolve_output_path(Path(cfg["testing"]["output_dir"]) / "musc_module_profile.xlsx")
    else:
        output_excel = resolve_output_path(output_path_arg)

    seed = 42
    model = MuScProfileRunner(cfg, seed=seed)

    profiler = ModuleProfiler(model.device)
    category_results = []
    start = time.perf_counter()
    for category in model.categories:
        print(f"Profiling category: {category}")
        category_results.append(
            profile_category(
                model=model,
                category=category,
                profiler=profiler,
                max_images=args.max_images,
            )
        )
    total_ms = (time.perf_counter() - start) * 1000

    output_excel = write_excel(output_excel, cfg, config_path, profiler, category_results, total_ms)
    print(f"MuSc module profile saved to: {output_excel}")


if __name__ == "__main__":
    main()
